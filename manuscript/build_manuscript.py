from pathlib import Path
import csv, json, re, html, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from paper_math import equation

OUT=Path(__file__).parent/'Naskah_Revisi';A=OUT/'audit'
def read(name):
    with (A/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
D={r['dataset']:r for r in read('audited_results.csv')};paths=json.loads((A/'selected_runs.json').read_text(encoding='utf-8'))
def num(r,k):return float(r[k])
def pct(r,k):return f'{100*num(r,k):.2f}'
def label(s):
    return s.replace('_efreq',' EF').replace('_ewidth',' EW').replace('wine_red','Red wine').replace('wine_white','White wine').replace('airquality','Air quality').replace('gasturbine','Gas turbine').replace('superconduct','Superconductivity').replace('california','California').replace('concrete','Concrete').replace('energy','Energy').replace('abalone','Abalone').replace('miami','Miami')
def model(s):return {'HistGradientBoosting':'HistGB','ExtraTrees':'Extra Trees','SVM_RBF':'RBF SVM','FTTransformer':'Custom FT','LogisticRegression':'Logistic','HistGBRegressor':'HistGB','CatBoostRegressor':'CatBoost','TabPFNRegressor':'TabPFN','ExtraTreesRegressor':'Extra Trees'}.get(s,s)
def table(headers,rows):return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,r))+' |' for r in rows])
order=['wine_red','wine_white','airquality','concrete_efreq','energy_efreq','abalone_efreq','california_efreq','miami_efreq','gasturbine_efreq','superconduct_efreq']
mainorder=['wine_red','wine_red_efreq','wine_white','wine_white_efreq','airquality','concrete_efreq','concrete_ewidth','energy_efreq','energy_ewidth','abalone_efreq','abalone_ewidth','california_efreq','california_ewidth','miami_efreq','miami_ewidth','gasturbine_efreq','gasturbine_ewidth','superconduct_efreq','superconduct_ewidth']
raw_n=[1599,4898,5000,1030,768,4177,20640,13776,36733,21263]
rows=[]
manifest=[]
for s,total in zip(order,raw_n):
    r=D[s];name=label(s).replace(' EF','');rows.append([name,f'{total:,}',f'{int(r["n"]):,}',r['p'],f'{int(r["n_test"]):,}'])
    au=json.loads((Path(paths[s])/'data_audit.json').read_text(encoding='utf-8'))
    manifest.append(dict(source=name,file=Path(au['source_path']).name,source_rows=total,analyzed_before_dedup=au['rows_raw'],analyzed_after_dedup=au['rows_clean'],features=au['features'],spec=au['spec'],source_sha256=au['source_sha256']))
(A/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
tables={'TABLE_DATA':table(['Source','Source rows','Analyzed n','Features','Test n'],rows)+'\n\nAnalyzed n is after subsampling and deduplication. EF denotes equal-frequency and EW equal-width targets. Test size is constant across each source\'s principal variants, although test membership can differ.'}
rows=[]
for r in read('protocol_comparison.csv'):
    rows.append([label(r['source']),f"{int(r['overlap'])}/{int(r['n_retained'])}",f"{num(r,'auc_dedup'):.4f}",f"{num(r,'auc_retained'):.4f}",f"{num(r,'delta_auc'):+.4f}"])
tables['TABLE_PROTOCOL']=table(['Source','Test overlap retained','AUC dedup','AUC retained','Difference'],rows)+'\n\nOverlap counts test rows whose predictor vector occurs in development. Difference is retained minus deduplicated. The test sets and selected primary models may differ across protocols.'
rows=[]
for r in read('ordinal_contrasts.csv'):
    rows.append([label(r['source']),f"{num(r,'ef_model'):.2f}",f"{num(r,'ef_constant'):.2f}",f"{num(r,'ew_model'):.2f}",f"{num(r,'ew_constant'):.2f}"])
tables['TABLE_ORDINAL']=table(['Source','EF model %','EF constant %','EW model %','EW constant %'],rows)
rows=[]
for s in mainorder:
    r=D[s];rows.append([label(s),model(r['classic']),f"{num(r,'classic_auc'):.4f}",model(r['neural']),f"{num(r,'neural_auc'):.4f}",f"{num(r,'tabpfn_auc'):.4f}"])
tables['TABLE_MODELS']=table(['Task','Classical','AUC','From scratch','AUC','TabPFN AUC'],rows)+'\n\nRepresentatives are selected by development ROC-AUC with the recorded tie-breakers. TabPFN is reported as a pretrained model category.'
B={r['dataset']:r for r in read('category_bootstrap.csv')};rows=[]
for s in mainorder:
    r=B[s]
    rows.append([label(s),f"{num(r,'scratch_delta'):+.4f} [{num(r,'scratch_lo'):+.4f}, {num(r,'scratch_hi'):+.4f}]",f"{num(r,'tabpfn_delta'):+.4f} [{num(r,'tabpfn_lo'):+.4f}, {num(r,'tabpfn_hi'):+.4f}]"])
tables['TABLE_INTERVALS']=table(['Task','From scratch minus classical','TabPFN minus classical'],rows)
rows=[]
for s in mainorder:
    r=D[s];rows.append([label(s),model(r['ordinal_model']),r['levels'],f"{num(r,'mae'):.4f}",pct(r,'exact'),pct(r,'w1'),f"{num(r,'mf1'):.4f}"])
tables['TABLE_ORDINAL_ALL']=table(['Task','Model','K','MAE','Exact %','Within 1 %','Macro-F1'],rows)
rows=[]
for r in read('ordinal_error_accounting.csv'):
    reduction=f"{num(r,'net_error_reduction_pct'):.2f}" if r['net_error_reduction_pct'] else '—'
    rows.append([label(r['dataset']),r['n'],r['baseline_errors'],r['corrected'],r['introduced'],r['model_errors'],reduction])
tables['TABLE_ERROR_ACCOUNTING']=table(['Task','n','Baseline errors','Corrected','Introduced','Model errors','Net reduction %'],rows)
rows=[]
for r in read('primary_runner_inference.csv'):
    effect=f"{num(r,'accuracy_delta_pp'):+.2f} [{num(r,'accuracy_lo_pp'):+.2f}, {num(r,'accuracy_hi_pp'):+.2f}]"
    rows.append([label(r['dataset']),r['n'],effect,f"{r['b']}/{r['c']}",f"{num(r,'mcnemar_p'):.4f}",f"{num(r,'holm_p'):.4f}"])
tables['TABLE_NULL_COMPARISONS']=table(['Task','n','Accuracy difference [95% CI]','b/c','Exact p','Holm p'],rows)
rows=[]
boundary={r['dataset']:r for r in read('rejection_boundary.csv')}
assert set(boundary)==set(mainorder),'rejection_boundary.csv must cover the 19 principal configurations'
for s in mainorder:
    r=boundary[s]
    rows.append([label(s),r['discordant'],r['observed_margin'],r['required_margin'] or '—',r['reached'],r['exact_p'],r['holm_p']])
tables['TABLE_RESOLUTION']=table(['Task','Discordant','Observed b−c','Required b−c','Reached','Exact p','Holm p'],rows)

text=(OUT/'audit'/'manuscript_template.md').read_text(encoding='utf-8')
refs=json.loads((OUT/'Reference'/'references.json').read_text(encoding='utf-8'));refsbykey={r['key']:r for r in refs};seen=[]
def cite(m):
    key=m[1];assert key in refsbykey
    if key not in seen:seen.append(key)
    return f'[{seen.index(key)+1}]'
text=re.sub(r'\{ref:([^}]+)\}',cite,text)
text=re.sub(r'(\[\d+\])\s+(?=\[\d+\])',r'\1, ',text)
assert len(seen)==len(refs)==22
for key,t in tables.items():text=text.replace('{{'+key+'}}',t)
eqs={'EQ1':'yᵢ = I(qᵢ ≥ 6),     (1)',
     'EQ2':'A₁ = (1/n) ∑ᵢ₌₁ⁿ I(|qᵢ − q̂ᵢ| ≤ 1),     (2)',
     'EQ3':'c* = arg max c∈Q [ (1/|D|) ∑ᵢ∈D I(|qᵢ − c| ≤ 1) ],     (3)',
     'EQ4':'G = Aₘ − A₀ = (R − H)/n = (E₀ − Eₘ)/n,     (4)',
     'EQ5':'η = (E₀ − Eₘ)/E₀ = G/(1 − A₀),   E₀ > 0.     (5)'}
for key,e in eqs.items():text=text.replace('{{'+key+'}}',e)
reftext='\n\n'.join(f'[{i}] {refsbykey[k]["ieee_text"]}' for i,k in enumerate(seen,1))
text=text.replace('{{REFERENCES}}',reftext)
assert '{{' not in text and '{ref:' not in text
(OUT/'Paper_Revised.md').write_text(text,encoding='utf-8')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.78);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
for style in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[style];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+attr),None)
    for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
    st.paragraph_format.space_after=Pt(6)
for st in doc.styles:
    for border in st.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True
for style,size in [('Heading 1',12),('Heading 2',11)]:
    st=doc.styles[style];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
doc.core_properties.title='Quantifying Predictive Gains in Tabular Learning through Ordinal Baselines and Protocol Analysis'
doc.core_properties.author='Heriyanto and Suharjito'
doc.core_properties.subject='Retrospective analysis of stored tabular learning experiments'
doc.core_properties.keywords='tabular learning; ordinal metrics; duplicate handling'

def add_inline(p,s):
    chunks=re.split(r'(\*\*.*?\*\*)',s)
    for part in chunks:
        r=p.add_run(part[2:-2] if part.startswith('**') else part)
        if part.startswith('**'):r.bold=True

def add_table(lines):
    rr=[[s.strip() for s in line.strip().strip('|').split('|')] for line in lines]
    rr=[rr[0]]+rr[2:];cols=len(rr[0]);t=doc.add_table(rows=1,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    compact=rr[0][1]=='Source rows'
    widths={3:[1.58,2.54,2.54],5:[1.62,1.26,1.26,1.26,1.26],6:[1.52,1.12,.7,1.10,.7,1.12],7:[1.52,1.02,.32,.80,.90,1.1,.60]}[cols]
    if rr[0][-1]=='Net reduction %':widths=[1.64,.40,.89,.84,.87,.87,1.01]
    if rr[0][-1]=='Holm p' and cols==6:widths=[1.58,.43,2.15,.49,.78,.78]
    if 'Required b−c' in rr[0]:widths=[1.66,.82,.98,.95,.68,.71,.72]
    assert len(widths)==cols,f'{cols} columns but {len(widths)} widths: {rr[0]}'
    for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for ri,row in enumerate(rr):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(.75 if compact else 3);p.paragraph_format.space_before=Pt(.75 if compact else 3);p.paragraph_format.line_spacing=1.0
            if compact:p.paragraph_format.keep_with_next=ri<len(rr)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or (cols in (6,7) and ci==1 and rr[0][-1] not in ['Net reduction %','Holm p']) else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(txt);r.font.size=Pt(8.6 if cols>=6 else 9.1);r.bold=ri==0
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for tag in ['top','left','bottom','right']:
                el=OxmlElement('w:'+tag);el.set(qn('w:w'),'30' if compact and tag in ['top','bottom'] else '75');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for tag in ['top','left','bottom','right']:
                el=OxmlElement('w:'+tag);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8E8E8');tcpr.append(shade)
        trpr=t.rows[ri]._tr.get_or_add_trPr();nonsplit=OxmlElement('w:cantSplit');trpr.append(nonsplit)
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=text.splitlines();i=0;in_refs=False;in_authors=True
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('|'):
        block=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i].strip());i+=1
        add_table(block);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;continue
    if line.startswith('## '):
        title=line[3:];in_refs=title=='References';p=doc.add_heading(title,level=1)
        if title=='Supplementary Results':p.paragraph_format.page_break_before=True
        continue
    if line.startswith('### '):
        p=doc.add_heading(line[4:],level=2)
        if line[4:] in ['B Ordinal outcomes for the principal configurations','D Development selection and final evaluation','C Paired tolerance errors for all principal tasks','D Paired binary comparison details']:p.paragraph_format.page_break_before=True
        continue
    if line.startswith('!['):
        source=re.search(r'\]\((.*?)\)',line)[1];p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        rr=p.add_run();rr.add_picture(str(OUT/source),width=Inches(6.0 if 'research_framework' in source else 6.4))
        pic=rr._r.xpath('.//wp:docPr')[0];pic.set('descr',Path(source).stem.replace('_',' '));continue
    if line.startswith('TABLE '):
        p=doc.add_paragraph(line,'Caption');p.paragraph_format.keep_with_next=True;continue
    if line.startswith('Fig. '):doc.add_paragraph(line,'Caption');continue
    if re.search(r'\([12345]\)$',line):
        equation(doc,int(re.search(r'\((\d)\)$',line)[1]));continue
    if line.startswith('**Abstract'):in_authors=False
    p=doc.add_paragraph()
    if in_refs:
        m=re.match(r'\[(\d+)\] (.*)',line);assert m,line
        ref=refsbykey[seen[int(m[1])-1]]
        p.add_run('['+m[1]+']\t')
        body=m[2];venue=ref.get('italic_venue','')
        if venue:
            before,after=body.split(venue,1);p.add_run(before);p.add_run(venue).italic=True;p.add_run(after)
        else:p.add_run(body)
        p.paragraph_format.left_indent=Inches(.30);p.paragraph_format.first_line_indent=Inches(-.30)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(.30));p.paragraph_format.keep_together=True
    else:add_inline(p,line)
    following=next((s.strip() for s in lines[i:] if s.strip()),'')
    if re.search(r'\([12345]\)$',following):p.paragraph_format.keep_with_next=True
    if in_refs:
        p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(5)
        for rr in p.runs:rr.font.size=Pt(9)
    elif in_authors:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for rr in p.runs:rr.font.size=Pt(9.5)
    else:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
dest=OUT/'Paper_Revised.docx';doc.save(dest)
abstract=text.split('**Abstract—** ')[1].split('**Index Terms')[0]
summary={'abstract_words':len(abstract.split()),'references':len(refs),'references_after_2020':sum(r['year']>2020 for r in refs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'word_count_approx':len(text.split()),'output_sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
(A/'document_checks.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
assert 150<=summary['abstract_words']<=250,summary
print(summary)

refnote=['# Reference manifest','', f'Daftar ini memuat {len(refs)} sumber yang dipakai dalam naskah; {sum(r["year"]>2020 for r in refs)} terbit setelah 2020. Sumber lama dipertahankan untuk metode asli dan data asal. Konferensi dan technical report tidak diberi label kuartil jurnal. Status kuartil yang telah diperiksa beserta tahun dan kategorinya dicatat dalam Catatan_Kuartil.md.','', 'Delapan PDF berasal dari koleksi pengguna. Metadata dua tambahan yang mendukung posisi kontribusi diperiksa melalui sumber penerbit dan Crossref pada 9 September 2026. Status berkas lokal dicatat per referensi di bawah. Sebagian PDF merupakan manuskrip penulis atau preprint.','']
for i,key in enumerate(seen,1):
    r=refsbykey[key];status=('PDF lokal: '+r['pdf']) if r.get('pdf') else 'PDF belum tersimpan lokal'
    refnote += [f'## {i} {key}',f'{r["text"]}',f'Sumber: {r["url"]}',f'Jenis: {r["type"]}; {status}.','']
(OUT/'Reference'/'README_Reference.md').write_text('\n'.join(refnote),encoding='utf-8')
