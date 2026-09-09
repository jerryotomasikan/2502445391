"""Editable OMML equations with explicit scripts, fractions and summation limits."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

def run(s,plain=False):
    r=OxmlElement('m:r')
    if plain:
        pr=OxmlElement('m:rPr');style=OxmlElement('m:sty');style.set(qn('m:val'),'p');pr.append(style);r.append(pr)
    t=OxmlElement('m:t');t.text=s;t.set(qn('xml:space'),'preserve');r.append(t);return r

def element(tag,items):
    e=OxmlElement(tag)
    for item in items:e.append(run(item) if isinstance(item,str) else item)
    return e

def sub(base,s):return element('m:sSub',[element('m:e',[base]),element('m:sub',[s])])
def sup(base,s):return element('m:sSup',[element('m:e',[base]),element('m:sup',[s])])
def frac(num,den):return element('m:f',[element('m:num',num),element('m:den',den)])
def hat(base):
    pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'̂');pr.append(ch)
    return element('m:acc',[pr,element('m:e',[base])])
def summation(lower,upper,expr):
    pr=OxmlElement('m:naryPr')
    for k,v in [('chr','∑'),('limLoc','subSup')]:
        e=OxmlElement('m:'+k);e.set(qn('m:val'),v);pr.append(e)
    if upper is None:
        hide=OxmlElement('m:supHide');hide.set(qn('m:val'),'1');pr.append(hide)
    return element('m:nary',[pr,element('m:sub',[lower]),element('m:sup',[] if upper is None else [upper]),element('m:e',expr)])
def indicator(*items):return [run('I',True),run('(',True),*items,run(')',True)]
def tolerance(i_pred):return indicator(run('|',True),sub('q','i'),run(' − ',True),i_pred,run('| ≤ 1',True))

def equation(doc,number):
    p=doc.add_paragraph();pf=p.paragraph_format
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.space_before=Pt(3);pf.space_after=Pt(6)
    pf.keep_together=True;pf.tab_stops.add_tab_stop(Inches(3.335),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Inches(6.67),WD_TAB_ALIGNMENT.RIGHT)
    math=OxmlElement('m:oMath')
    if number==1:items=[sub('y','i'),run(' = ',True),*indicator(sub('q','i'),run(' ≥ 6',True)),run(',',True)]
    elif number==2:items=[sub('A','1'),run(' = ',True),frac([run('1',True)],['n']),summation('i = 1','n',tolerance(sub(hat('q'),'i'))),run(',',True)]
    elif number==3:
        items=[sup('c','*'),run(' = ',True),sub(run('arg max',True),'c ∈ Q'),run(' ',True),frac([run('1',True)],['|D|']),summation('i ∈ D',None,tolerance('c')),run(',',True)]
    elif number==4:
        items=['G',run(' = ',True),sub('A','m'),run(' − ',True),sub('A','0'),run(' = ',True),frac(['R',run(' − ',True),'H'],['n']),run(' = ',True),frac([sub('E','0'),run(' − ',True),sub('E','m')],['n']),run(',',True)]
    elif number==5:
        items=['η',run(' = ',True),frac([sub('E','0'),run(' − ',True),sub('E','m')],[sub('E','0')]),run(' = ',True),frac(['G'],[run('1 − ',True),sub('A','0')]),run(',   ',True),sub('E','0'),run(' > 0.',True)]
    else:raise ValueError(number)
    for item in items:math.append(run(item) if isinstance(item,str) else item)
    p.add_run('\t');p._p.append(math);p.add_run('\t('+str(number)+')')
    return p
