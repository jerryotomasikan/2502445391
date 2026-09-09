from pathlib import Path
import json,re,csv
import numpy as np,pandas as pd
from sklearn.metrics import cohen_kappa_score,f1_score,roc_auc_score
from scipy.stats import rankdata,binomtest
import sys
# python-docx and pypdf come from the project venv (.venv-build); the earlier
# borrow from the Codex runtime cache is not durable across tool updates.
from docx import Document
from pypdf import PdfReader

O=Path(__file__).resolve().parent;A=O.parent/'results';P=json.loads((A/'selected_runs.json').read_text());D=pd.read_csv(A/'audited_results.csv').set_index('dataset');checks=[]
def ok(v,label):
    checks.append(dict(check=label,passed=bool(v)))
    if not v:raise AssertionError(label)
for s,path in P.items():
    p=Path(path);op=pd.read_csv(p/'test_predictions_ordinal.csv');tm=pd.read_csv(p/'test_metrics_ordinal.csv').set_index('model');au=json.loads((p/'data_audit.json').read_text());by=pd.read_csv(p/'test_predictions_binary.csv')
    for m,r in tm.iterrows():
        ok(np.isclose(f1_score(op.quality_true,op['pred_'+m],labels=au['ordinal_levels'],average='macro',zero_division=0),r.macro_f1),f'{s}/{m} ordinal macro-F1')
        ok(np.isclose(cohen_kappa_score(op.quality_true,op['pred_'+m],weights='quadratic'),r.qwk),f'{s}/{m} QWK')
    # Validate rank-based AUC bootstrap computation against sklearn on real paired resamples.
    y=by.binary_true.to_numpy();neg=np.flatnonzero(y==0);pos=np.flatnonzero(y==1);rng=np.random.default_rng(123)
    for m in [D.loc[s,'classic'],D.loc[s,'neural'],'TabPFN']:
        prob=np.clip(by['p_'+m].to_numpy(),1e-7,1-1e-7)
        for k in range(5):
            ix=np.r_[rng.choice(neg,len(neg)),rng.choice(pos,len(pos))];rs=rankdata(prob[ix]);v=(rs[len(neg):].sum()-len(pos)*(len(pos)+1)/2)/(len(pos)*len(neg))
            ok(np.isclose(v,roc_auc_score(y[ix],prob[ix]),atol=1e-13),f'{s}/{m} bootstrap AUC formula sample{k}')
main=D[~D.index.str.contains('__seed|nodedup')]
ok(len(main)==19,'19 principal configurations')
ok((main.neural_auc>main.classic_auc).sum()==5,'scratch category count')
ok((main.tabpfn_auc>main.classic_auc).sum()==15,'TabPFN category count')
ps=sorted(main.mcn_p);adj=np.minimum(1,np.maximum.accumulate(np.array(ps)*np.arange(19,0,-1)))
ok(round(adj[0],4)==.2245,'Holm adjusted minimum p')
ok(round(binomtest(5,5,.5).pvalue,4)==.0625,'conditional five-discordance example')
md=(O/'Paper_Revised.md').read_text(encoding='utf-8');doc=Document(O/'Paper_Revised.docx');pdf=PdfReader(O/'Paper_Revised.pdf')
body=md.split('## References')[0];citations=[int(x) for x in re.findall(r'\[(\d+)\]',body)]
ok(list(dict.fromkeys(citations))==list(range(1,23)),'citation first appearance order 1..22')
ok(len(doc.tables)==9,'nine Word tables (I-IV, S1-S5)')
ok(len(doc.inline_shapes)==4,'four Word figures including the research framework')
for i,block in enumerate(re.findall(r'(?m)^\|.*(?:\n\|.*)*',md)):
    values=[[c.strip() for c in r.strip('|').split('|')] for r in block.splitlines()];values=[values[0]]+values[2:]
    actual=[[c.text for c in r.cells] for r in doc.tables[i].rows]
    ok(actual==values,f'Word table {i+1} exactly matches generated source')
pdftext='\n'.join(p.extract_text() for p in pdf.pages)
for bad in ['{{','{ref:','\ufffd','\u25a0','\u25a1','\u00e2\u20ac','\u00c3\u009f']:
    ok(bad not in md and bad not in pdftext,f'no placeholder or missing-glyph marker {repr(bad)}')
abstract=md.split('**Abstract—** ')[1].split('**Index Terms')[0]
ok(150<=len(abstract.split())<=250,'abstract within 150-250 words')
ok(len(pdf.pages)==len(json.loads((A/'pdf_render'/'pages.json').read_text())), 'PDF page count matches rendered-page manifest')
ok(re.findall(r'(?m)^Fig\. (\d+)\.',md)==['1','2','3','4'],'figure captions numbered 1-4 in order')
ok('Figure_1_research_framework.png' in md,'framework figure embedded in source')
baseline=pd.read_csv(A/'ordinal_baseline_uncertainty.csv').set_index('dataset')
for name,successes,constant_successes,lo,hi in [('wine_red',263,256,.74,4.78),('miami_ewidth',975,957,1.02,2.76)]:
    r=baseline.loc[name]
    ok(r.model_successes==successes and r.constant_successes==constant_successes,name+' tolerance counts')
    ok(round(r.ci_lo*100,2)==lo and round(r.ci_hi*100,2)==hi,name+' stated interval endpoints')
    ok(np.isclose(r.difference,(r.only_model-r.only_constant)/r.n),name+' paired tolerance difference')
(A/'final_validation.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
print('ADDITIONAL CHECKS PASSED',len(checks))
